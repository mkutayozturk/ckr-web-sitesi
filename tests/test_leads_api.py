import os
import uuid
from io import BytesIO

import pytest
from openpyxl import load_workbook
from pymongo import MongoClient


MONGO_TEST_URL = os.environ.get("MONGO_TEST_URL")


def require_mongo():
    if not MONGO_TEST_URL:
        pytest.skip("MONGO_TEST_URL tanımlı değil; gerçek MongoDB testi atlandı.")
    client = MongoClient(MONGO_TEST_URL, serverSelectionTimeoutMS=1200)
    try:
        client.server_info()
    except Exception as exc:
        pytest.skip(f"Gerçek MongoDB bağlantısı yok: {exc}")
    return client


@pytest.fixture()
def api_client():
    mongo_client = require_mongo()
    db_name = f"ckr_test_{uuid.uuid4().hex}"
    os.environ["MONGO_URL"] = MONGO_TEST_URL
    os.environ["DB_NAME"] = db_name
    os.environ["ADMIN_PASSWORD"] = "test-admin"
    os.environ["ADMIN_JWT_SECRET"] = "test-secret"

    from backend import server
    from fastapi.testclient import TestClient

    with TestClient(server.app) as client:
        yield client

    mongo_client.drop_database(db_name)


def lead_payload(phone="0530 111 22 33"):
    return {
        "request_type": "seller",
        "property_type": "Daire",
        "location": {
            "mode": "single",
            "province": "Çanakkale",
            "province_code": "17",
            "district": "Merkez",
            "district_code": "1230",
            "districts": ["Merkez"],
            "neighborhood": "Barbaros",
            "settlement": "Barbaros",
            "settlement_type": "neighborhood",
            "settlement_code": "mah-20098",
            "raw": "Çanakkale / Merkez / Barbaros Mahallesi",
            "selections": [
                {
                    "scope": "settlement",
                    "province": "Çanakkale",
                    "province_code": "17",
                    "district": "Merkez",
                    "district_code": "1230",
                    "settlement": "Barbaros",
                    "settlement_code": "mah-20098",
                    "settlement_type": "neighborhood",
                    "label": "Barbaros Mahallesi",
                }
            ],
        },
        "property_features": {"metrekare": "120", "oda": "3+1"},
        "price_expectation": "8.000.000 TL",
        "timing": "Yüksek",
        "contact": {"name": "Test Müşteri", "phone": phone, "email": "test@example.com"},
        "source": {"form_type": "satici", "button_source": "test"},
        "motivation": {"answers": {"aciliyet": 80}, "score": 80, "category": "Yüksek", "summary": "Test özeti"},
        "kvkk_notice_presented": True,
        "kvkk_notice_presented_at": "2026-10-08T12:01:00+00:00",
        "kvkk_notice_version": "2026-10-12",
        "kvkk_acknowledged": False,
        "marketing_consent": False,
        "started_at": "2026-10-08T12:00:00+00:00",
    }


def test_lead_create_admin_update_and_export(api_client):
    created = api_client.post("/api/leads", json=lead_payload())
    assert created.status_code == 201
    lead = created.json()
    assert lead["follow_up_status"] == "new"
    assert lead["admin_notes"] == "Takip edilecek yarım başvuru"
    assert lead["motivation"]["score"] == 80

    duplicate = api_client.post("/api/leads", json=lead_payload())
    assert duplicate.status_code == 409

    preference = api_client.patch(f"/api/leads/{lead['id']}/preference", json={"preference": "ADVISOR_REQUESTED"})
    assert preference.status_code == 200
    assert preference.json()["follow_up_status"] == "qualified"
    assert preference.json()["admin_notes"] == "Formu danışmana yolla tercihi seçildi."

    unauthorized = api_client.get("/api/admin/leads")
    assert unauthorized.status_code == 401

    login = api_client.post("/api/admin/login", json={"password": "test-admin"})
    assert login.status_code == 200
    token = login.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    listed = api_client.get("/api/admin/leads", headers=headers)
    assert listed.status_code == 200
    assert listed.json()["total"] == 1

    filtered = api_client.get("/api/admin/leads?province=Çanakkale&district=Merkez&settlement=Barbaros", headers=headers)
    assert filtered.status_code == 200
    assert filtered.json()["total"] == 1

    updated = api_client.patch(
        f"/api/admin/leads/{lead['id']}",
        headers=headers,
        json={"follow_up_status": "contacted", "admin_notes": "Arandı, tekrar görüşülecek."},
    )
    assert updated.status_code == 200
    assert updated.json()["follow_up_status"] == "contacted"
    assert updated.json()["admin_notes"] == "Arandı, tekrar görüşülecek."

    exported = api_client.get("/api/admin/leads/export.xlsx", headers=headers)
    assert exported.status_code == 200
    assert exported.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    workbook = load_workbook(BytesIO(exported.content))
    assert "Müşteri Başvuruları" in workbook.sheetnames
    sheet = workbook["Müşteri Başvuruları"]
    headers_row = [cell.value for cell in sheet[1]]
    assert "Ad soyad" in headers_row
    assert "İl" in headers_row
    assert "Çoklu bölge tercihleri" in headers_row
    assert "Başvuru durumu" in headers_row
    assert "KVKK aydınlatma sunuldu" in headers_row
    name_column = headers_row.index("Ad soyad") + 1
    assert sheet.cell(row=2, column=name_column).value == "Test Müşteri"


def test_draft_autosave_complete_and_admin_filters(api_client):
    payload = lead_payload(phone="0530 111 22 44")
    payload["application_id"] = f"app-{uuid.uuid4().hex}"
    payload["motivation"] = {
        "answers": {"completion": "incomplete", "responses": {}, "missing": ["Soru"]},
        "score": None,
        "category": "Değerlendirilmedi",
        "summary": "Motivasyon soruları eksik yanıtlandığı için motivasyon puanı hesaplanmadı.",
    }

    draft = api_client.post("/api/leads/draft", json=payload)
    assert draft.status_code == 201
    draft_body = draft.json()
    assert draft_body["application_status"] == "draft"
    assert draft_body["kvkk_notice_presented"] is True
    assert draft_body["kvkk_acknowledged"] is False
    assert draft_body["draft_token"]
    assert draft_body["follow_up_status"] == "follow_up"

    missing_notice_payload = lead_payload(phone="0530 111 22 55")
    missing_notice_payload["application_id"] = f"app-{uuid.uuid4().hex}"
    missing_notice_payload["kvkk_notice_presented"] = False
    missing_notice_payload["kvkk_acknowledged"] = False
    missing_notice = api_client.post("/api/leads/draft", json=missing_notice_payload)
    assert missing_notice.status_code == 400

    payload["property_features"]["metrekare"] = "125"
    forbidden = api_client.patch(f"/api/leads/{payload['application_id']}/draft", json=payload)
    assert forbidden.status_code == 403

    headers = {"X-CKR-Draft-Token": draft_body["draft_token"]}
    autosaved = api_client.patch(f"/api/leads/{payload['application_id']}/draft", headers=headers, json=payload)
    assert autosaved.status_code == 200
    assert autosaved.json()["property_features"]["metrekare"] == "125"
    assert autosaved.json()["application_status"] == "draft"

    login = api_client.post("/api/admin/login", json={"password": "test-admin"})
    token = login.json()["access_token"]
    admin_headers = {"Authorization": f"Bearer {token}"}
    draft_list = api_client.get("/api/admin/leads?application_status=draft", headers=admin_headers)
    assert draft_list.status_code == 200
    assert draft_list.json()["total"] == 1

    payload["motivation"] = {"answers": {"aciliyet": 80}, "score": 80, "category": "Yüksek", "summary": "Test özeti"}
    completed = api_client.patch(f"/api/leads/{payload['application_id']}/complete", headers=headers, json=payload)
    assert completed.status_code == 200
    completed_body = completed.json()
    assert completed_body["id"] == draft_body["id"]
    assert completed_body["application_status"] == "completed"
    assert completed_body["follow_up_status"] == "new"
    assert completed_body["draft_expires_at"] is None

    completed_again = api_client.patch(f"/api/leads/{payload['application_id']}/complete", headers=headers, json=payload)
    assert completed_again.status_code == 200
    assert completed_again.json()["id"] == draft_body["id"]

    completed_list = api_client.get("/api/admin/leads?application_status=completed", headers=admin_headers)
    assert completed_list.status_code == 200
    assert completed_list.json()["total"] == 1
